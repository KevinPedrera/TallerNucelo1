from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "entrega"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "Informe_Taller5_Nucleo1_ITSQMET.docx"

BLUE = "1F4E79"
LIGHT = "D9EAF7"
GRAY = "F2F2F2"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table, color="D9D9D9"):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders")
        tblPr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            b.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def set_cell(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, "FFFFFF")
        shade(t.rows[0].cells[i], BLUE)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val)
            if r_i % 2 == 1:
                shade(cells[i], GRAY)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Taller 5 Núcleo 1 ITSQMET | Informe académico").font.size = Pt(8)

# Cover with explicit markers because the referenced CARATULA.docx was not available locally.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("[MARCADOR: LOGOTIPO ITSQMET]")
r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)

for text, size, bold in [
    ("INSTITUTO SUPERIOR TECNOLÓGICO", 15, True),
    ("TALLER INTEGRADOR NÚCLEO 1", 18, True),
    ("Informe del Taller 5 Caso Integrador", 16, True),
    ("Gestión de usuarios productos y pedidos", 12, False),
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text); run.font.size = Pt(size); run.bold = bold

doc.add_paragraph().paragraph_format.space_after = Pt(120)
for text in ["Estudiante: [MARCADOR: NOMBRE DEL ESTUDIANTE]", "Carrera: Desarrollo de Software", "Docente: [MARCADOR: NOMBRE DEL DOCENTE]", "Fecha: 9 de octubre de 2026"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text)
doc.add_page_break()

doc.add_heading("Introducción", level=1)
doc.add_paragraph("Este informe documenta la implementación del Taller 5 del Núcleo 1 de ITSQMET. El proyecto integra Spring Boot, Spring MVC, Spring Data JPA, MySQL, Thymeleaf, una API REST y Spring Security para resolver un caso de gestión de usuarios, productos, categorías y pedidos.")
doc.add_paragraph("La solución fue construida de forma incremental sobre el proyecto existente. Las pruebas automatizadas se ejecutaron con H2 en memoria y la instancia demo se inició correctamente en el puerto 8080. La conexión final a MySQL queda parametrizada por variables de entorno y requiere validar las credenciales del equipo de entrega.")

doc.add_heading("Problemática", level=1)
doc.add_paragraph("La guía oficial plantea una empresa que actualmente utiliza hojas de cálculo y registros manuales. Este mecanismo favorece la duplicidad de información, dificulta la trazabilidad y no ofrece un control adecuado de accesos. El sistema implementado centraliza la información y aplica permisos diferenciados para las operaciones administrativas.")

doc.add_heading("Objetivos", level=1)
doc.add_heading("Objetivo general", level=2)
doc.add_paragraph("Implementar una aplicación web segura y persistente para administrar usuarios, productos, categorías y pedidos mediante una arquitectura MVC basada en Spring Boot.")
doc.add_heading("Objetivos específicos", level=2)
add_bullets(doc, ["Modelar las entidades y relaciones con JPA.", "Implementar operaciones CRUD para productos y categorías.", "Registrar usuarios con contraseñas cifradas mediante BCrypt.", "Exponer una API REST con autorización por método HTTP y rol.", "Validar el funcionamiento mediante pruebas automatizadas y pruebas HTTP."])

doc.add_heading("Alcance del proyecto", level=1)
add_bullets(doc, ["Registro e inicio de sesión de usuarios.", "Roles USER y ADMIN.", "Gestión de categorías y productos.", "Búsqueda y cálculo de unidades y valor del inventario.", "Registro y consulta de pedidos.", "Endpoints REST para productos, categorías y pedidos.", "Vistas Thymeleaf con Bootstrap."])

doc.add_heading("Tecnologías utilizadas", level=1)
table(doc, ["Tecnología", "Uso"], [
    ["Java 21", "Lenguaje de desarrollo"], ["Spring Boot 4.1.1", "Configuración y ejecución"],
    ["Spring MVC y Thymeleaf", "Interfaz web"], ["Spring Data JPA e Hibernate", "Persistencia"],
    ["MySQL", "Base de datos de entrega"], ["H2", "Pruebas automatizadas y perfil demo"],
    ["Spring Security", "Autenticación y autorización"], ["Maven Wrapper", "Construcción y pruebas"],
])

doc.add_heading("Arquitectura MVC", level=1)
doc.add_paragraph("La aplicación se organiza en capas. Los controladores reciben las solicitudes, los servicios concentran las reglas de negocio, los repositorios acceden a la base de datos y las entidades representan el modelo persistente. Las vistas Thymeleaf conforman la capa de presentación y los controladores REST exponen recursos JSON.")
table(doc, ["Capa", "Implementación"], [
    ["Model", "Usuario, Categoria, Producto, Pedido, Role y EstadoPedido"],
    ["Repository", "Repositorios JpaRepository para las cuatro entidades principales"],
    ["Service", "Reglas de registro, CRUD, búsquedas y cálculos de inventario"],
    ["Controller", "MVC web, autenticación y controladores REST"],
    ["View", "Plantillas login, registro, productos y pedidos"],
])

doc.add_heading("Estructura del proyecto", level=1)
doc.add_paragraph("La estructura relevante se encuentra dentro de src/main/java/com/example/tallernucleo1:")
add_bullets(doc, ["config/SecurityConfig.java", "model/", "repository/", "service/", "controller/", "src/main/resources/templates/", "src/test/java/", "src/test/resources/application-test.properties"])

doc.add_heading("Modelo de datos y relaciones JPA", level=1)
doc.add_paragraph("Usuario se relaciona con Pedido mediante una relación uno a muchos. Categoria se relaciona con Producto mediante una relación uno a muchos. Pedido y Producto se relacionan mediante una relación muchos a muchos con la tabla intermedia pedido_producto.")
table(doc, ["Entidad", "Campos principales", "Relación"], [
    ["Usuario", "id, username, password, nombreCompleto, role", "1:N con Pedido"],
    ["Categoria", "id, nombre", "1:N con Producto"],
    ["Producto", "id, nombre, descripcion, precio, stock", "N:1 con Categoria y N:M con Pedido"],
    ["Pedido", "id, fecha, direccionEntrega, estado", "N:1 con Usuario y N:M con Producto"],
])

doc.add_heading("Entidades repositorios servicios y controladores", level=1)
doc.add_paragraph("Los repositorios extienden JpaRepository y usan consultas derivadas para búsquedas. ProductoService resuelve la categoría, actualiza productos y calcula el valor total del inventario. PedidoService valida los productos seleccionados y asocia el pedido al usuario autenticado. Los controladores separan la navegación web de los endpoints REST.")

doc.add_heading("Interfaz web", level=1)
doc.add_paragraph("La interfaz incluye inicio de sesión, registro, listado y búsqueda de productos, formulario de alta y edición, eliminación protegida por sesión, creación de pedidos y consulta de pedidos. Bootstrap se utiliza para la estructura visual y la adaptación básica a diferentes tamaños de pantalla.")
doc.add_paragraph("[MARCADOR DE EVIDENCIA: insertar captura de la página principal y del formulario de producto]")

doc.add_heading("API REST", level=1)
table(doc, ["Método", "Ruta", "Acceso"], [
    ["GET", "/api/productos y /api/productos/{id}", "Público"],
    ["POST", "/api/productos", "ADMIN"], ["PUT", "/api/productos/{id}", "ADMIN"],
    ["DELETE", "/api/productos/{id}", "ADMIN"], ["GET", "/api/categorias", "Público"],
    ["POST", "/api/categorias", "ADMIN"], ["GET/POST", "/api/pedidos", "Usuario autenticado"],
])

doc.add_heading("Spring Security", level=1)
doc.add_paragraph("La configuración mantiene CSRF activo para la aplicación web y lo ignora únicamente para /api/**, donde las pruebas utilizan HTTP Basic. Las rutas GET públicas de productos y categorías no requieren autenticación. Las operaciones POST, PUT y DELETE administrativas requieren ROLE_ADMIN. Las contraseñas se codifican con BCrypt y el usuario ADMIN inicial se crea desde variables de entorno.")

doc.add_heading("Pruebas realizadas y resultados", level=1)
table(doc, ["Funcionalidad", "Procedimiento", "Esperado", "Obtenido", "Estado"], [
    ["Compilación", ".\\mvnw.cmd test", "Compilar sin errores", "Compilación exitosa", "OK"],
    ["Contexto Spring", "TallerNucleo1ApplicationTests", "Cargar contexto", "Contexto H2 cargado", "OK"],
    ["GET productos", "GET /api/productos", "200 público", "200 con producto registrado", "OK"],
    ["POST sin credenciales", "POST /api/productos", "401", "401", "OK"],
    ["POST ADMIN", "POST /api/categorias con Basic Auth", "201", "201", "OK"],
    ["POST USER", "POST /api/categorias con usuario USER", "403", "403", "OK"],
    ["Inicio demo", "spring-boot:run -Dspring-boot.run.profiles=demo", "Aplicación iniciada", "Tomcat 8080 iniciado", "OK"],
    ["MySQL", "Arranque con configuración predeterminada", "Conectar a BD", "Pendiente de credenciales MySQL", "PENDIENTE"],
])

doc.add_paragraph("Las pruebas automatizadas corresponden a ApiSecurityTests y TallerNucleo1ApplicationTests. Las pruebas de API adicionales se ejecutaron contra la instancia demo local mediante solicitudes HTTP.")

doc.add_heading("Evidencias visuales", level=1)
doc.add_paragraph("Se verificaron visualmente, mediante navegador, las siguientes pantallas: inicio de sesión con confirmación de registro, página principal de productos vacía y página de productos con el producto de prueba, stock 10 y valor de inventario 459.00.")
for item in ["[CAPTURA PENDIENTE: estructura del proyecto en IDE]", "[CAPTURA PENDIENTE: tablas y relaciones en MySQL]", "[CAPTURA PENDIENTE: ejecución de pruebas Maven]", "[CAPTURA PENDIENTE: respuesta REST en Postman con 200, 201, 401, 403 y 404]"]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph("No se incluyen imágenes simuladas. Las capturas pendientes deben obtenerse en el equipo de entrega con MySQL activo y Postman/IDE abiertos.")

doc.add_heading("Conclusiones", level=1)
doc.add_paragraph("La implementación cubre la arquitectura MVC, persistencia JPA, CRUD de productos y categorías, gestión de pedidos, autenticación, autorización por roles, BCrypt, API REST y pruebas automatizadas. La instancia demo funciona y permite demostrar los flujos principales sin depender de una instalación externa.")
doc.add_paragraph("La validación final contra MySQL y la incorporación de la carátula institucional permanecen pendientes porque no se dispuso de las credenciales de base de datos ni del archivo CARATULA.docx referenciado.")

doc.add_heading("Recomendaciones", level=1)
add_bullets(doc, ["Configurar DB_USERNAME y DB_PASSWORD sin escribirlos en archivos versionados.", "Tomar las capturas pendientes con la aplicación conectada a MySQL.", "Reemplazar los marcadores de carátula antes de entregar el informe.", "Cambiar ADMIN_PASSWORD para la demostración final."])

doc.add_heading("Referencias documentales", level=1)
doc.add_paragraph("Guía de la Unidad de Integración Curricular Integración de la Información, Taller 5 Caso Práctico.")
doc.add_paragraph("Transcripciones N1_Clase_01_05Oct, N1_Clase_02_06Oct, N1_Clase_03_07Oct y N1_Clase_04_09Oct.")

doc.save(DOCX)
print(DOCX)
